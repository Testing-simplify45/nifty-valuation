"""Writes results to a Google Sheet (if set up) or an Excel file for download."""
import io
import numpy as np
import pandas as pd


def _clean(df):
    out = df.copy()
    for c in out.columns:
        if pd.api.types.is_float_dtype(out[c]):
            out[c] = out[c].round(3)
    out = out.replace({np.nan: ""})
    return out.astype(object).where(out.notna(), "")


def to_excel_bytes(master, summary, log, all_options=None):
    from openpyxl.styles import PatternFill, Font
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        sheets = [("Master", master)]
        if all_options is not None:
            sheets.append(("All Options", all_options))
        sheets += [("Expiry Summary", summary), ("Log", log)]
        for name, df in sheets:
            _clean(df).to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            if name in ("Master", "All Options") and "Status" in df.columns:
                col = list(df.columns).index("Status") + 1
                for r in range(2, ws.max_row + 1):
                    v = ws.cell(r, col).value
                    fill = {"Overvalued": "F8C9C9", "Undervalued": "C9E8C9", "Low confidence": "E6E6E6"}.get(v)
                    if fill:
                        ws.cell(r, col).fill = PatternFill("solid", start_color=fill)
            for i, c in enumerate(df.columns, 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = max(11, min(22, len(str(c)) + 3))
    return buf.getvalue()


def write_gsheet(master, summary, log_new, sheet_id, service_account_info, all_options=None):
    """Replaces Master and Expiry Summary; appends new rows to Log."""
    import gspread
    gc = gspread.service_account_from_dict(service_account_info)
    sh = gc.open_by_key(sheet_id)

    def ws_for(title):
        try:
            return sh.worksheet(title)
        except gspread.WorksheetNotFound:
            return sh.add_worksheet(title=title, rows=1000, cols=40)

    tabs = [("Master", master)]
    if all_options is not None:
        tabs.append(("All Options", all_options))
    tabs.append(("Expiry Summary", summary))
    for title, df in tabs:
        ws = ws_for(title)
        ws.clear()
        c = _clean(df)
        ws.update([list(c.columns)] + c.values.tolist(), value_input_option="USER_ENTERED")
        ws.freeze(rows=1)
        ws.format("1:1", {"textFormat": {"bold": True}})

    ws = ws_for("Log")
    c = _clean(log_new)
    if not ws.get_all_values():
        ws.update([list(c.columns)])
    if len(c):
        ws.append_rows(c.values.tolist(), value_input_option="USER_ENTERED")
    return sh.url


def to_matrix(df):
    """Header + rows as plain Python lists, ready for the Sheets API."""
    c = _clean(df)
    return [list(c.columns)] + c.values.tolist()
